# -*- coding: utf-8 -*-
"""快速统计每月河鱼/海鱼可刷出数量"""
import collections
import openpyxl

SRC = r'C:/Users/Administrator/Documents/WPSDrive/344975910/WPS云盘/动森鱼虫合集.xlsx'
SEA = {'大海', '栈桥', '所有海'}

wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)

wa = wb['鱼分析表']
rows_wa = list(wa.iter_rows(min_row=3, values_only=True))
months = collections.defaultdict(set)
for row in rows_wa:
    iid = row[0]
    if iid is None:
        continue
    iid = int(iid)
    m = row[5]
    w = row[7]
    if w and w > 0 and m:
        months[iid].add(int(m))

ws = wb['鱼结论记录']
rows_ws = list(ws.iter_rows(min_row=2, values_only=True))
riv = collections.Counter()
sea = collections.Counter()
for row in rows_ws:
    iid = row[0]
    if iid is None:
        continue
    iid = int(iid)
    at = str(row[5])
    tgt = sea if at in SEA else riv
    for m in months.get(iid, set()):
        tgt[m] += 1

print('月份  河鱼  海鱼')
for m in range(1, 13):
    print(f'{m:2d}月  {riv[m]:4d}  {sea[m]:4d}')

n_riv = sum(1 for r in rows_ws if r[0] is not None and str(r[5]) not in SEA)
n_sea = sum(1 for r in rows_ws if r[0] is not None and str(r[5]) in SEA)
print('河鱼鱼种数', n_riv, ' 海鱼鱼种数', n_sea)

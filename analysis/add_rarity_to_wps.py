# -*- coding: utf-8 -*-
"""
add_rarity_to_wps.py
把 fish.json 的 4 档稀有度写入《动森鱼虫合集.xlsx》「鱼结论记录」页签的 H 列。

对齐链路：表内 ItemID -> FishStatusParam.DebugName(日文) -> fish.json.name-JPja -> name-CNzh -> rarity
垃圾道具（空罐/长靴/轮胎/石头）在 fish.json 中不存在，留空。

注意：G 列「月份总数」是 XLOOKUP 公式，故 H 列追加静态值；同步把自动筛选扩展到 H。
"""
import openpyxl, json, csv, io
from pathlib import Path
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(r"C:/Users/Administrator/WorkBuddy/acnh")
XLSX = Path(r"C:/Users/Administrator/Documents/WPSDrive/344975910/WPS云盘/动森鱼虫合集.xlsx")
SHEET = "鱼结论记录"
FISH_JSON = Path(r"C:/Users/Administrator/Desktop/fish.json")

# ---- ItemID -> 稀有度 ----
rows = list(csv.reader(io.open(ROOT / "data/FishStatusParam.csv", encoding="utf-8")))
hdr, data = rows[0], [r for r in rows[1:] if any(r)]
ITEM, DBG = hdr.index("ItemID"), hdr.index("DebugName")
jp = {int(r[ITEM]): r[DBG].strip("'") for r in data}

fj = json.load(io.open(FISH_JSON, encoding="utf-8"))
jp2zh = {v["name"]["name-JPja"]: v["name"]["name-CNzh"] for v in fj.values()}
zh2rar = {v["name"]["name-CNzh"]: v["availability"]["rarity"] for v in fj.values()}
item2rar = {i: zh2rar[jp2zh[j]] for i, j in jp.items() if j in jp2zh}

# ---- 写入 ----
wb = openpyxl.load_workbook(XLSX)
ws = wb[SHEET]

HCOL = 8
ws.cell(row=1, column=HCOL, value="稀有度")
ws.cell(row=1, column=HCOL).font = Font(name="宋体", size=11, bold=True)
ws.cell(row=1, column=HCOL).alignment = Alignment(horizontal="left", vertical="top")

filled = blank = 0
detail = []
for r in range(2, ws.max_row + 1):
    v = ws.cell(row=r, column=1).value
    if v is None or not str(v).strip():
        continue
    try:
        item = int(v)
    except (TypeError, ValueError):
        continue
    rar = item2rar.get(item)
    c = ws.cell(row=r, column=HCOL)
    c.font = Font(name="宋体", size=11)
    c.alignment = Alignment(horizontal="left")
    if rar:
        c.value = rar
        filled += 1
    else:
        blank += 1
    detail.append((r, item, ws.cell(row=r, column=2).value, rar))

ws.column_dimensions["H"].width = 13
ref = ws.auto_filter.ref
if ref:
    ws.auto_filter.ref = "A1:H89"
print(f"自动筛选: {ref} -> {ws.auto_filter.ref}")

wb.save(XLSX)
print(f"已写入: 有稀有度 {filled} 行 / 留空 {blank} 行 / 合计 {filled+blank} 行")
print()
print("=== 写入明细（前 12 + 后 8）===")
for d in detail[:12]:
    print(f"  r{d[0]:<3} {d[1]:>6} {str(d[2]):<10} -> {d[3] or '(留空·垃圾)'}")
print("  ...")
for d in detail[-8:]:
    print(f"  r{d[0]:<3} {d[1]:>6} {str(d[2]):<10} -> {d[3] or '(留空·垃圾)'}")

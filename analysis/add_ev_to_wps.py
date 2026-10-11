# -*- coding: utf-8 -*-
"""
add_ev_to_wps.py
在《动森鱼虫合集》新增页签「期望收益」，写入：
  ① 河鱼：每月 × 4 钓点（河流/池塘/悬崖上/出海口）期望收益 + 河系均值
  ② 海鱼：每月 × 2 钓点（大海/栈桥）期望收益 + 海系均值
  ③ 河海对比摘要

数据来源：analysis/河鱼期望收益.csv / 海鱼期望收益.csv（由 expected_value.py 生成）。
口径：E = Σ(Weight×Price)/ΣWeight；三时段(白天/晨昏/夜晚)等长取平均。

安全策略：写入同目录临时文件后 os.replace 原子替换，原文件不被写坏。
"""
import csv, io, os, tempfile, shutil, datetime
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

SRC = r'C:/Users/Administrator/Documents/WPSDrive/344975910/WPS云盘/动森鱼虫合集.xlsx'
BAK = r'C:/Users/Administrator/WorkBuddy/acnh/backup'
ROOT = r'C:/Users/Administrator/WorkBuddy/acnh/analysis'
SHEET = '期望收益'

# ---- 样式 ----
F_TITLE = Font(name='微软雅黑', size=14, bold=True, color='FF1F4E9C')
F_SEC   = Font(name='微软雅黑', size=12, bold=True, color='FFFFFFFF')
F_HDR   = Font(name='宋体', size=11, bold=True, color='FF1F2A37')
F_BODY  = Font(name='宋体', size=11)
F_NOTE  = Font(name='宋体', size=9, italic=True, color='FF6B7280')
FILL_SEC  = PatternFill('solid', fgColor='FF1F4E9C')
FILL_HDR  = PatternFill('solid', fgColor='FFDCE6F5')
FILL_ALT  = PatternFill('solid', fgColor='FFF5F8FC')
FILL_SUM  = PatternFill('solid', fgColor='FFFFF3D6')
CENTER = Alignment(horizontal='center', vertical='center')
LEFT   = Alignment(horizontal='left', vertical='center')
THIN = Side(style='thin', color='FFBFC8D6')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NUMFMT = '#,##0.0'


def read_csv(path):
    with open(path, 'r', encoding='utf-8-sig', newline='') as f:
        rows = list(csv.reader(f))
    return rows[0], rows[1:]


def put(ws, r, c, v, font=F_BODY, fill=None, align=CENTER, fmt=None, border=True):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font
    if fill:
        cell.fill = fill
    cell.alignment = align
    if fmt:
        cell.number_format = fmt
    if border:
        cell.border = BORDER
    return cell


def main():
    head_r, body_r = read_csv(os.path.join(ROOT, '河鱼期望收益.csv'))
    head_s, body_s = read_csv(os.path.join(ROOT, '海鱼期望收益.csv'))

    wb = openpyxl.load_workbook(SRC)
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET, index=9)   # 插到「水域表」之后、「Sheet1」之前

    nc = 6
    # ---- 标题 ----
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=nc)
    put(ws, 1, 1, 'ACNH 钓鱼 · 每月每钓点期望收益（单位：铃钱／次）',
        font=F_TITLE, align=LEFT, border=False)
    ws.row_dimensions[1].height = 24

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=nc)
    put(ws, 2, 1, '口径：期望收益 E = Σ(Weight×Price) / ΣWeight；一天三时段等长（各约 8h），月值取三档平均。'
                  '钓点鱼池由「水域表」决定（河流/悬崖上/出海口共用河流鱼，栈桥 ⊇ 大海）。',
        font=F_NOTE, align=LEFT, border=False)

    # ---- 河鱼 ----
    r = 4
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=nc)
    c = put(ws, r, 1, '河鱼（4 钓点）', font=F_SEC, fill=FILL_SEC, align=LEFT)
    for cc in range(2, nc + 1):
        put(ws, r, cc, None, fill=FILL_SEC)
    r += 1
    hdr_row = r
    for i, h in enumerate(head_r):
        put(ws, r, i + 1, h, font=F_HDR, fill=FILL_HDR)
    r += 1
    for ri, row in enumerate(body_r):
        fill = FILL_ALT if ri % 2 else None
        put(ws, r, 1, row[0], font=F_HDR, fill=fill)                       # 月份
        for i, v in enumerate(row[1:], start=2):
            is_sum = (i == nc)                                             # 均值列
            put(ws, r, i, float(v), fill=(FILL_SUM if is_sum else fill),
                font=(F_HDR if is_sum else F_BODY), fmt=NUMFMT)
        r += 1

    r += 1
    # ---- 海鱼 ----
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=nc)
    put(ws, r, 1, '海鱼（2 钓点）', font=F_SEC, fill=FILL_SEC, align=LEFT)
    for cc in range(2, nc + 1):
        put(ws, r, cc, None, fill=FILL_SEC)
    r += 1
    for i, h in enumerate(head_s):
        put(ws, r, i + 1, h, font=F_HDR, fill=FILL_HDR)
    r += 1
    for ri, row in enumerate(body_s):
        fill = FILL_ALT if ri % 2 else None
        put(ws, r, 1, row[0], font=F_HDR, fill=fill)
        for i, v in enumerate(row[1:], start=2):
            is_sum = (i == len(head_s))
            put(ws, r, i, float(v), fill=(FILL_SUM if is_sum else fill),
                font=(F_HDR if is_sum else F_BODY), fmt=NUMFMT)
        r += 1

    r += 1
    # ---- 河海对比 ----
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=nc)
    put(ws, r, 1, '河海对比（同口径）', font=F_SEC, fill=FILL_SEC, align=LEFT)
    for cc in range(2, nc + 1):
        put(ws, r, cc, None, fill=FILL_SEC)
    r += 1
    cmp_hdr = ['指标', '河鱼（4 钓点均值）', '海鱼（2 钓点均值）']
    for i, h in enumerate(cmp_hdr):
        put(ws, r, i + 1, h, font=F_HDR, fill=FILL_HDR)
    r += 1
    cmp_rows = [
        ('收益区间', '443 ~ 1457', '1199 ~ 1694'),
        ('极差', '1014', '495'),
        ('相对波动 CV', '≈ 38%', '≈ 13%'),
        ('最低月', '4 月  443', '4 月  1199'),
        ('最高月', '9 月  1457', '9 月  1694'),
    ]
    for ri, row in enumerate(cmp_rows):
        fill = FILL_ALT if ri % 2 else None
        for i, v in enumerate(row):
            put(ws, r, i + 1, v, font=(F_HDR if i == 0 else F_BODY), fill=fill,
                align=(LEFT if i == 0 else CENTER))
        r += 1

    r += 1
    put(ws, r, 1, '结论：海钓 = 稳定高收益（低方差）；河钓 = 季节性爆发（高方差）。'
                  '海鱼最低月（4 月 1199）已超河鱼全年均值；叠加「池塘」逆周期钓点（12 月 1075 为当月第一）。',
        font=F_NOTE, align=LEFT, border=False)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=nc)

    # ---- 列宽 ----
    for col, w in zip('ABCDEF', [10, 13, 13, 13, 13, 13]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A6'
    ws.sheet_view.showGridLines = False

    # ---- 强制打开时重算（保护原有公式列的显示值） ----
    try:
        wb.calculation.fullCalcOnLoad = True
    except Exception as e:
        print('warning: set fullCalcOnLoad failed:', e)

    # ---- 备份 + 原子替换 ----
    os.makedirs(BAK, exist_ok=True)
    ts = datetime.datetime.now().strftime('%Y%m%d_%H%M')
    bak = os.path.join(BAK, f'动森鱼虫合集_备份_{ts}.xlsx')
    shutil.copy2(SRC, bak)

    d = os.path.dirname(SRC)
    fd, tmp = tempfile.mkstemp(suffix='.xlsx', dir=d)
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, SRC)
    except Exception:
        if os.path.exists(tmp):
            os.remove(tmp)
        raise

    print('备份:', bak)
    print('已写入 SRC:', SRC)
    print(f'页签「{SHEET}」: {ws.max_row} x {ws.max_column}')
    print('sheets:', wb.sheetnames)


if __name__ == '__main__':
    main()

# -*- coding: utf-8 -*-
"""
add_prob_to_wps.py
把「鱼分析表」页签 V/W 列（刷鱼概率平均值 / 最大值）搬到「鱼结论记录」页签。

结构：
  鱼分析表：88 条鱼 × 36 行/鱼 = 3168 行（行 3..3170），每块首行有 V=平均值 W=最大值
  鱼结论记录：88 行数据（行 2..89），A=ItemID
  对齐键：ItemID（鱼分析表 A 列 == 鱼结论记录 A 列）

写入：I 列「刷鱼概率均值」、J 列「刷鱼概率最大值」
"""
import openpyxl, shutil
from pathlib import Path
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

SRC = Path(r"C:/Users/Administrator/Documents/WPSDrive/344975910/WPS云盘/动森鱼虫合集.xlsx")
ROOT = Path(r"C:/Users/Administrator/WorkBuddy/acnh")
OUT = ROOT / "tmp_动森鱼虫合集_v2.xlsx"
BLOCK = 36

wb = openpyxl.load_workbook(SRC, data_only=True)

# ---- 1. 从鱼分析表收集每鱼 (V, W) ----
ws_a = wb["鱼分析表"]
prob = {}
for start in range(3, ws_a.max_row + 1, BLOCK):
    a = ws_a.cell(row=start, column=1).value
    if a is None:
        continue
    v = ws_a.cell(row=start, column=22).value
    w = ws_a.cell(row=start, column=23).value
    prob[int(a)] = (v, w)
print(f"鱼分析表: 收集到 {len(prob)} 条鱼的 概率均值/最大值")

# ---- 2. 写入鱼结论记录 I/J 列 ----
wb2 = openpyxl.load_workbook(SRC)
ws = wb2["鱼结论记录"]

for col, title in [(9, "刷鱼概率均值"), (10, "刷鱼概率最大值")]:
    c = ws.cell(row=1, column=col, value=title)
    c.font = Font(name="宋体", size=11, bold=True)
    c.alignment = Alignment(horizontal="left", vertical="top")

miss, filled = [], 0
detail = []
for r in range(2, ws.max_row + 1):
    idv = ws.cell(row=r, column=1).value
    if idv is None or not str(idv).strip():
        continue
    item = int(idv)
    v, w = prob.get(item, (None, None))
    for col, val in [(9, v), (10, w)]:
        c = ws.cell(row=r, column=col)
        c.font = Font(name="宋体", size=11)
        c.alignment = Alignment(horizontal="left")
        if val is not None:
            c.value = round(float(val), 6)
    if v is None:
        miss.append((r, item))
    else:
        filled += 1
    detail.append((r, item, ws.cell(row=r, column=2).value, v, w))

ws.column_dimensions["I"].width = 13
ws.column_dimensions["J"].width = 15
ws.auto_filter.ref = "A1:J89"
print(f"写入: 有值 {filled} 行 / 缺失 {len(miss)} 行")
if miss:
    print("  缺失:", miss)

wb2.save(OUT)
print(f"已写: {OUT}")
print()
print("=== 抽样核对 ===")
for d in detail[:6] + detail[10:14] + detail[36:40] + detail[-5:]:
    vs = f"{d[3]:.4%}" if d[3] is not None else "—"
    wsv = f"{d[4]:.4%}" if d[4] is not None else "—"
    print(f"  r{d[0]:<3} {d[1]:>6} {str(d[2]):<10} 均值={vs:<9} 最大={wsv}")

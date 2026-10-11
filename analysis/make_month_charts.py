# -*- coding: utf-8 -*-
"""
make_month_charts.py
生成「鱼类出没月份图」，两种排序 × 河/海 = 4 张：
  - 河鱼月份图_月份数优先.png
  - 海鱼月份图_月份数优先.png
  - 河鱼月份图_时段优先.png
  - 海鱼月份图_时段优先.png

形式：横向时间轴矩阵（Gantt 风格）
  纵轴 = 每种鱼，两种排序：
      A「月份数优先」：① 可出没月份数↓ ② 活动时间优先级 ③ ItemID↑
      B「时段优先」  ：① 活动时间优先级 ② 可出没月份数↓ ③ ItemID↑
  横轴 = 1–12 月
  填充 = 该鱼在该月权重 > 0（色块月份之间无间隙，连续呈现）
  行右侧 = Size 徽章（9 种）+ 售价 + 活动时间
  底部汇总带 = 该月可刷出的鱼种总数

配色（4 色系，相近色系内保持区分度）：
  全天            -> 蓝色系
  非白天 & 夜间    -> 紫色系（深紫 / 浅紫）
  非夜间 & 白天    -> 橙黄系（橙 / 金黄）
  非晨昏          -> 绿色系

钓点标记（按 AppearType）：
  池塘   -> 黑色粗边框
  悬崖上 -> 斜线纹理
  栈桥   -> 斜线纹理（与悬崖上一致）
  出海口 -> 交叉网格纹理
  其余（河流/所有河/大海/所有海）-> 无标记
"""
import collections
import openpyxl
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import Patch
from matplotlib.font_manager import FontProperties

matplotlib.rcParams['hatch.linewidth'] = 1.1

SRC = r'C:/Users/Administrator/Documents/WPSDrive/344975910/WPS云盘/动森鱼虫合集.xlsx'

OUT = {
    ('river', 'month'): r'C:/Users/Administrator/WorkBuddy/acnh/analysis/河鱼月份图_月份数优先.png',
    ('sea',   'month'): r'C:/Users/Administrator/WorkBuddy/acnh/analysis/海鱼月份图_月份数优先.png',
    ('river', 'time'):  r'C:/Users/Administrator/WorkBuddy/acnh/analysis/河鱼月份图_时段优先.png',
    ('sea',   'time'):  r'C:/Users/Administrator/WorkBuddy/acnh/analysis/海鱼月份图_时段优先.png',
}

# 中文字体
CN = FontProperties(fname=r'C:/Windows/Fonts/msyh.ttc')

# 4 色系（相近色系，组内拉开明度/色相差）
COLOR = {
    '全天':   '#2E5FA3',   # 蓝色系
    '非白天': '#6B3FA0',   # 紫色系 - 深紫
    '夜间':   '#B39DDB',   # 紫色系 - 浅紫
    '非夜间': '#C96A1E',   # 橙黄系 - 橙
    '白天':   '#F0BE4A',   # 橙黄系 - 金黄
    '非晨昏': '#2F8F4E',   # 绿色系
}
# 活动时间优先级（用户指定）
ORDER = ['全天', '白天', '非夜间', '非白天', '夜间', '非晨昏']

SEA_TYPES = {'大海', '栈桥', '所有海'}

# 钓点纹理
HATCH = {
    '悬崖上': '///',
    '栈桥':   '///',
    '出海口': 'xxx',
}
POND = '池塘'

# Size 徽章（9 种：SS/S/M/L/LL/LLL/J/K/U）
SIZE_BADGE_BG = '#4A5260'
SIZE_BADGE_EC = '#2F3640'
try:
    MONO = FontProperties(fname=r'C:/Windows/Fonts/consola.ttf', size=8)
except Exception:
    MONO = FontProperties(size=8)


def _lum(hexc):
    """感知亮度"""
    r, g, b = (int(hexc[i:i + 2], 16) for i in (1, 3, 5))
    return 0.299 * r + 0.587 * g + 0.114 * b


def contrast(hexc):
    """在给定底色上取对比色（用于纹理线）"""
    return '#1A1A1A' if _lum(hexc) > 140 else '#FFFFFF'


def load():
    wb = openpyxl.load_workbook(SRC)
    ws = wb['鱼结论记录']
    wa = wb['鱼分析表']

    months = collections.defaultdict(set)
    for r in range(3, wa.max_row + 1):
        iid = wa.cell(row=r, column=1).value
        if iid is None:
            continue
        iid = int(iid)
        m = wa.cell(row=r, column=6).value
        w = wa.cell(row=r, column=8).value
        if w and w > 0 and m:
            months[iid].add(int(m))

    rows = []
    for r in range(2, ws.max_row + 1):
        iid = ws.cell(row=r, column=1).value
        if iid is None:
            continue
        iid = int(iid)
        rows.append({
            'iid': iid,
            'name': ws.cell(row=r, column=2).value,
            'size': str(ws.cell(row=r, column=3).value),
            'price': ws.cell(row=r, column=4).value or 0,
            'at': str(ws.cell(row=r, column=6).value),
            'time': ws.cell(row=r, column=7).value,
            'side': 'sea' if str(ws.cell(row=r, column=6).value) in SEA_TYPES else 'river',
            'months': months.get(iid, set()),
        })
    return rows


def sort_month(rows):
    """排序 A：月份数↓ → 活动时间 → ItemID"""
    return sorted(rows, key=lambda x: (-len(x['months']),
                                       ORDER.index(x['time']) if x['time'] in ORDER else 99,
                                       x['iid']))


def sort_time(rows):
    """排序 B：活动时间 → 月份数↓ → ItemID"""
    return sorted(rows, key=lambda x: (ORDER.index(x['time']) if x['time'] in ORDER else 99,
                                       -len(x['months']),
                                       x['iid']))


def draw(rows, title, outpath, side_cn):
    n = len(rows)
    fig_h = max(6, n * 0.34 + 1.6)
    fig, ax = plt.subplots(figsize=(13.6, fig_h))
    fig.patch.set_facecolor('#FFFFFF')

    x_min, x_max = 0.5, 12.5

    # ---- 每月汇总带 ----
    SUM_ROW = -2.2
    SUM_H = 1.0
    y_lo = SUM_ROW - SUM_H / 2 - 0.35
    y_hi = n - 0.5
    ax.add_patch(plt.Rectangle((x_min, SUM_ROW - SUM_H / 2),
                               x_max - x_min, SUM_H,
                               facecolor='#EDEFF3', edgecolor='none', zorder=0.5))

    # 月份交替底纹
    for m in range(1, 13):
        ax.axvspan(m - 0.5, m + 0.5,
                   ymin=0, ymax=1, facecolor='#FAFAF8' if m % 2 == 0 else '#FFFFFF',
                   zorder=0)

    used = set()
    used_at = set()
    for yi, row in enumerate(rows):
        y = n - 1 - yi
        t, mset, at = row['time'], row['months'], row['at']
        if not mset:
            continue
        used.add(t)
        used_at.add(at)
        fc = COLOR.get(t, '#999999')
        ms = sorted(mset)
        spans = []
        seg, prev = ms[0], ms[0]
        for m in ms[1:]:
            if m == prev + 1:
                prev = m
            else:
                spans.append((seg, prev))
                seg = prev = m
        spans.append((seg, prev))

        # 钓点标记样式
        kw = dict(facecolor=fc, zorder=2)
        if at == POND:
            kw.update(edgecolor='#111111', linewidth=1.8)
        elif at in HATCH:
            kw.update(hatch=HATCH[at], edgecolor=contrast(fc), linewidth=0.0)
        else:
            kw.update(edgecolor='none', linewidth=0.0)

        for a, b in spans:
            ax.add_patch(plt.Rectangle((a - 0.5, y - 0.42),
                                       (b - a + 1.0), 0.84, **kw))

    # 每月总数
    monthly = collections.Counter()
    for row in rows:
        for m in row['months']:
            monthly[m] += 1

    ax.set_xlim(x_min, x_max)
    ax.set_ylim(y_lo, y_hi)
    ax.set_xticks(range(1, 13))
    ax.set_xticklabels([f'{m}月' for m in range(1, 13)], fontproperties=CN, fontsize=10)
    ax.set_yticks(range(n))
    ax.set_yticklabels([r['name'] for r in reversed(rows)],
                       fontproperties=CN, fontsize=8.5)
    ax.xaxis.set_ticks_position('top')
    ax.xaxis.set_label_position('top')
    ax.tick_params(axis='x', length=0, pad=4)
    ax.tick_params(axis='y', length=0, pad=3)

    # 汇总带数值
    mx = max(monthly.values()) if monthly else 1
    for m in range(1, 13):
        v = monthly.get(m, 0)
        alpha = 0.12 + 0.55 * (v / mx)
        ax.add_patch(plt.Rectangle((m - 0.5, SUM_ROW - SUM_H / 2),
                                   1.0, SUM_H,
                                   facecolor=COLOR['全天'], alpha=alpha,
                                   edgecolor='none', zorder=1.2))
        ax.text(m, SUM_ROW, str(v), ha='center', va='center',
                fontproperties=CN, fontsize=9.5, color='#12233F',
                fontweight='bold', zorder=3)

    for sp in ax.spines.values():
        sp.set_visible(False)
    for m in range(1, 13):
        ax.axvline(m + 0.5, color='#E2E2E2', linewidth=0.7, zorder=1, alpha=0.9)
    ax.axvline(0.5, color='#E2E2E2', linewidth=0.7, zorder=1, alpha=0.9)
    ax.set_axisbelow(True)

    ax.axhline(SUM_ROW + SUM_H / 2 + 0.1, color='#C9CDD6', linewidth=0.9, zorder=1.5)
    ax.text(-1.15, SUM_ROW, f'{side_cn}鱼合计', fontproperties=CN, fontsize=9,
            va='center', ha='right', color='#5A6270', fontweight='bold')

    # ---- 行右侧：Size 徽章 / 售价 / 活动时间 ----
    for yi, row in enumerate(rows):
        y = n - 1 - yi
        t = row['time']
        # Size 徽章（9 种：SS/S/M/L/LL/LLL/J/K/U）
        ax.text(12.62, y, row['size'], fontproperties=MONO, fontsize=8,
                va='center', ha='center', color='white', zorder=4,
                bbox=dict(boxstyle='round,pad=0.38', facecolor=SIZE_BADGE_BG,
                          edgecolor=SIZE_BADGE_EC, linewidth=0.6))
        # 售价（右对齐，千分位）
        ax.text(14.02, y, f"{int(row['price']):,}", fontproperties=MONO, fontsize=8,
                va='center', ha='right', color='#2A2F38', zorder=4)
        # 活动时间
        ax.text(14.28, y, t, fontproperties=CN, fontsize=7.5,
                va='center', ha='left', color=COLOR.get(t, '#666'))

    ax.set_title(title, fontproperties=CN, fontsize=14, pad=26, loc='left')

    # 图例
    legend_order = [k for k in ORDER if k in used]
    handles = [Patch(facecolor=COLOR[k], edgecolor='none', label=k) for k in legend_order]
    leg1 = ax.legend(handles=handles, loc='lower center', bbox_to_anchor=(0.5, -0.038),
                     ncol=max(1, len(legend_order)), frameon=False, prop=CN, fontsize=9)
    ax.add_artist(leg1)

    # 钓点标记图例：只显示本图实际出现的钓点
    marks = []
    if POND in used_at:
        marks.append(('池塘（黑框）', dict(facecolor='#C9CDD6', edgecolor='#111111', linewidth=1.8)))
    diag = [s for s in ('悬崖上', '栈桥') if s in used_at]
    if diag:
        marks.append((' / '.join(diag) + '（斜线）',
                      dict(facecolor='#C9CDD6', edgecolor='#333333', hatch='///', linewidth=0.0)))
    if '出海口' in used_at:
        marks.append(('出海口（网格）',
                      dict(facecolor='#C9CDD6', edgecolor='#333333', hatch='xxx', linewidth=0.0)))
    if marks:
        spot_handles = [Patch(label=lb, **kw) for lb, kw in marks]
        ax.legend(handles=spot_handles, loc='lower center', bbox_to_anchor=(0.5, -0.075),
                  ncol=len(spot_handles), frameon=False, prop=CN, fontsize=8.5)

    ax.annotate('纵轴：鱼种　横轴：月份　色块：该月可刷出　行右：Size / 售价(铃钱) / 活动时间　'
                '底部汇总带：该月可刷出的鱼种总数',
                xy=(0, 0), xytext=(0, -0.108), xycoords='axes fraction',
                textcoords='axes fraction', fontproperties=CN, fontsize=9,
                color='#777777', va='top')

    plt.tight_layout(rect=[0, 0.075, 0.88, 0.98])
    fig.savefig(outpath, dpi=170, facecolor='#FFFFFF', bbox_inches='tight')
    plt.close(fig)
    print(f'saved: {outpath}  ({n} 行)')


def main():
    rows = load()
    river = [r for r in rows if r['side'] == 'river']
    sea = [r for r in rows if r['side'] == 'sea']
    print(f'河: {len(river)}  海: {len(sea)}')

    draw(sort_month(river), '河鱼 · 出没月份图（月份数优先 · 北半球）',
         OUT[('river', 'month')], '河')
    draw(sort_month(sea), '海鱼 · 出没月份图（月份数优先 · 北半球）',
         OUT[('sea', 'month')], '海')
    draw(sort_time(river), '河鱼 · 出没月份图（时段优先 · 北半球）',
         OUT[('river', 'time')], '河')
    draw(sort_time(sea), '海鱼 · 出没月份图（时段优先 · 北半球）',
         OUT[('sea', 'time')], '海')


if __name__ == '__main__':
    main()
